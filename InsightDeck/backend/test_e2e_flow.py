import io
import requests
import pptx

from pathlib import Path

sample_path = Path("audit_sales_demo.xlsx")
if not sample_path.exists():
    sample_path = Path(__file__).resolve().parent.parent / "audit_sales_demo.xlsx"

print('=== 1. UPLOAD EXCEL ===')
with open(sample_path, 'rb') as f:
    r_up = requests.post(
        'http://localhost:8000/upload',
        files={'file': ('audit_sales_demo.xlsx', f, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    )
print('Upload status:', r_up.status_code)
up_data = r_up.json()
print('Filename:', up_data.get('filename'))
print('Details:', up_data.get('details'))

print('\n=== 2. GET PRESENTATION MODEL ===')
r_pres = requests.get('http://localhost:8000/presentation')
print('Pres status:', r_pres.status_code)
pres_data = r_pres.json()
slides = pres_data.get('slides', [])
print('Slide count:', len(slides))
for idx, s in enumerate(slides):
    print(f"  Slide {idx + 1}: type={s.get('type')}, title={s.get('title')}")

print('\n=== 3. EDIT SLIDE 1 & APPLY CHANGES ===')
slides[0]['title'] = 'Audit Executive Board Review'
slides[0]['subtitle'] = 'Q3 Compliance & Fiscal Review'
slides[0]['bullets'] = [
    'Verified by Internal Audit Committee',
    'Scope: 12 high-value enterprise transactions',
    '100% Deterministic Accounting Verification'
]
slides[0]['content'] = slides[0]['bullets']

print('\n=== 4. GENERATE PPTX VIA POST /generate-presentation ===')
r_gen = requests.post('http://localhost:8000/generate-presentation', json={'presentation': pres_data})
print('Response status:', r_gen.status_code)
print('Content-Type:', r_gen.headers.get('content-type'))
print('Binary PPTX size:', len(r_gen.content), 'bytes')
assert r_gen.status_code == 200
assert 'presentationml.presentation' in r_gen.headers.get('content-type')
assert len(r_gen.content) > 5000

print('\n=== 5. VALIDATE PPTX INTEGRITY WITH PYTHON-PPTX ===')
prs = pptx.Presentation(io.BytesIO(r_gen.content))
print('Parsed slides in PPTX:', len(prs.slides))
assert len(prs.slides) == 8

# Slide 1 text verification
s1_texts = [shape.text_frame.text for shape in prs.slides[0].shapes if shape.has_text_frame]
s1_combined = ' '.join(s1_texts)
print('Slide 1 text sample:', s1_combined[:150])
assert 'Audit Executive Board Review' in s1_combined
assert 'Q3 Compliance & Fiscal Review' in s1_combined

# Slide 4 Chart verification
s4_charts = [shape for shape in prs.slides[3].shapes if shape.has_chart]
print('Slide 4 charts count:', len(s4_charts))
assert len(s4_charts) == 1
chart = s4_charts[0].chart
print('Slide 4 chart type:', chart.chart_type)

# Save to disk
with open('outputs/test_verified_presentation.pptx', 'wb') as f_out:
    f_out.write(r_gen.content)
print('Successfully saved outputs/test_verified_presentation.pptx')
print('\n=== ALL END-TO-END VALIDATIONS PASSED! ===')
